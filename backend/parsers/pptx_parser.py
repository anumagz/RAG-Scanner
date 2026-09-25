from pathlib import Path

from pptx import Presentation

from parsers.base import BaseParser, ParsedDocument


class PPTXParser(BaseParser):

    def parse(
        self,
        file_path: str,
        file_id: int | None = None,
    ):

        path = Path(file_path)

        presentation = Presentation(
            str(path)
        )

        documents = []

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):

            texts = []

            for shape in slide.shapes:

                if not hasattr(
                    shape,
                    "text",
                ):
                    continue

                if shape.text.strip():

                    texts.append(
                        shape.text.strip()
                    )

            text = "\n".join(texts)

            if not text:
                continue

            documents.append(
                ParsedDocument(
                    text=text,
                    metadata={
                        "file": path.name,
                        "file_path": str(path),
                        "file_id": file_id,
                        "slide": slide_number,
                        "type": "pptx",
                    },
                )
            )

        return documents
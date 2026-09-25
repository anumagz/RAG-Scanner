from pathlib import Path

from openpyxl import load_workbook

from parsers.base import BaseParser, ParsedDocument


class XLSXParser(BaseParser):

    def parse(
        self,
        file_path: str,
        file_id: int | None = None,
    ):

        path = Path(file_path)

        workbook = load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True,
        )

        documents = []

        try:

            for sheet in workbook.worksheets:

                rows = []

                for row in sheet.iter_rows(
                    values_only=True
                ):

                    values = []

                    for value in row:

                        if value is None:
                            values.append("")
                        else:
                            values.append(
                                str(value)
                            )

                    if any(values):

                        rows.append(
                            " | ".join(values)
                        )

                text = "\n".join(rows)

                if not text.strip():
                    continue

                documents.append(
                    ParsedDocument(
                        text=text,
                        metadata={
                            "file": path.name,
                            "file_path": str(path),
                            "file_id": file_id,
                            "sheet": sheet.title,
                            "type": "xlsx",
                        },
                    )
                )

        finally:
            workbook.close()

        return documents
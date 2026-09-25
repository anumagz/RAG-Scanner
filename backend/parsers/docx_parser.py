from pathlib import Path

from docx import Document

from parsers.base import BaseParser, ParsedDocument


class DOCXParser(BaseParser):

    def parse(
        self,
        file_path: str,
        file_id: int | None = None,
    ):

        path = Path(file_path)

        document = Document(
            str(path)
        )

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        text = "\n".join(paragraphs)

        return [
            ParsedDocument(
                text=text,
                metadata={
                    "file": path.name,
                    "file_path": str(path),
                    "file_id": file_id,
                    "type": "docx",
                },
            )
        ]